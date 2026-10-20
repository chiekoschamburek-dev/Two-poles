# -*- coding: utf-8 -*-
import re
from docx import Document
from lxml import etree
from docx.oxml.ns import qn

p = r'D:\桌面\project\project\两仪·TwoPoles_项目需求分析与概要设计文档_修订版.docx'
d = Document(p)
print('paragraphs', len(d.paragraphs), 'tables', len(d.tables), 'inline', len(d.inline_shapes))
print('core title:', d.core_properties.title)
print('core author:', d.core_properties.author)
print('core keywords:', d.core_properties.keywords)
print('--- 6.1.1 ---')
for i in range(154, 161):
    x = d.paragraphs[i]
    print(i, x.style.name, 'pb', x.paragraph_format.page_break_before, 'keep', x.paragraph_format.keep_with_next,
          'ind', x.paragraph_format.first_line_indent, repr(x.text[:50]))
print('--- level3 sample ---')
for i in (55, 70, 79, 156, 160, 164, 167):
    x = d.paragraphs[i]
    ppr = x._p.find(qn('w:pPr'))
    print(i, x.style.name, repr(x.text[:30]), 'bold', [r.bold for r in x.runs], 'size', [r.font.size.pt if r.font.size else None for r in x.runs], 'color', [r.font.color.rgb if r.font.color and r.font.color.type else None for r in x.runs])
    if ppr is not None:
        print('   pPr children:', [etree.QName(c).localname for c in ppr])
print('--- images ---')
for sh in d.inline_shapes:
    blip = sh._inline.find('.//' + qn('a:blip'))
    rid = blip.get(qn('r:embed'))
    print(rid, d.part.rels[rid].target_ref, f'{sh.width.cm:.2f}x{sh.height.cm:.2f}')
print('--- table1 rows ---')
t = d.tables[1]
for ri, row in enumerate(t.rows):
    print(ri, [c.text[:30] for c in row.cells], 'keep', [p.paragraph_format.keep_with_next for c in row.cells for p in c.paragraphs])
print('--- target text ---')
for i in (3, 6, 71, 159, 175, 177, 180):
    print(i, repr(d.paragraphs[i].text[:180]))
print('--- table 2 r1/r5 ---')
print(repr(d.tables[2].cell(1,2).text))
print(repr(d.tables[2].cell(5,2).text))
print('--- table 5 F11 ---')
print(repr(d.tables[5].cell(11,2).text))
print('--- table 7 M3 ---')
print(repr(d.tables[7].cell(3,2).text))
print('--- table 9 row7 ---')
print(repr(d.tables[9].cell(7,0).text))
print('--- 其它/0至10 ---')
full = '\n'.join(x.text for x in d.paragraphs)
print('其它', full.count('其它'), '0至10', full.count('0至10'))
print('F01～F20', full.count('F01～F20'))
