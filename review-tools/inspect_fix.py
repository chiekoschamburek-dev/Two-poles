# -*- coding: utf-8 -*-
import re, sys
from docx import Document
from docx.shared import Cm
from lxml import etree

src = r'D:\桌面\project\project\artifacts\fitness-design\两仪·TwoPoles_项目需求分析与概要设计文档.docx'
doc = Document(src)
ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}

print('=== P70 pPr ===')
print(etree.tostring(doc.paragraphs[70]._p.find('w:pPr', ns), pretty_print=True, encoding='unicode'))

print('=== level3 direct pPr summary ===')
for i, p in enumerate(doc.paragraphs):
    if re.match(r'^\d+\.\d+\.\d+\s', p.text.strip()):
        ppr = p._p.find('w:pPr', ns)
        bits = []
        if ppr is not None:
            for child in ppr:
                tag = etree.QName(child).localname
                if tag in ('pStyle', 'ind', 'spacing', 'jc'):
                    bits.append(etree.tostring(child, encoding='unicode'))
        print(i, p.text.strip()[:30], p.style.name, ' | '.join(bits))
        if i > 80: break

print('=== inline shapes ===')
for idx, sh in enumerate(doc.inline_shapes):
    inline = sh._inline
    blip = inline.find('.//a:blip', ns)
    rid = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed') if blip is not None else None
    target = doc.part.rels[rid].target_ref if rid else None
    print(idx, rid, target, f'{sh.width.cm:.2f}x{sh.height.cm:.2f} cm', f'aspect={sh.width/sh.height:.3f}')

print('=== table 1 info ===')
t = doc.tables[1]
print('rows', len(t.rows), 'cols', len(t.columns))
for ri, row in enumerate(t.rows):
    trpr = row._tr.find('w:trPr', ns)
    print('row', ri, 'trPr=', etree.tostring(trpr, encoding='unicode') if trpr is not None else None)
    print('   ', [c.text[:20] for c in row.cells])
