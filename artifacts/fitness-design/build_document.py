from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
import hashlib, json, math, os
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(r'D:\桌面\project\project\artifacts\fitness-design')
# 课程模板 docx 的路径改由环境变量提供：原先这里硬编码了微信接收目录（含个人账号标识）
REF_ENV=os.environ.get('TWOPOLES_TEMPLATE_DOCX')
if not REF_ENV: raise SystemExit('请设置 TWOPOLES_TEMPLATE_DOCX 指向课程模板 docx 的本地路径')
REF=Path(REF_ENV)
ASSET=ROOT/'assets'; ASSET.mkdir(exist_ok=True)
OUT=ROOT/'两仪·TwoPoles_项目需求分析与概要设计文档.docx'
FONT=r'C:\Windows\Fonts\msyh.ttc'; BOLD=r'C:\Windows\Fonts\msyhbd.ttc'
NAVY='#16334b'; TEAL='#087e80'; BLUE='#2468a0'; MUTED='#627585'; BG='#f2f6f9'; LINE='#dce5ed'
def font(n=26,b=False): return ImageFont.truetype(BOLD if b else FONT,n)
def text(d,xy,s,n=26,c=NAVY,b=False): d.text(xy,s,font=font(n,b),fill=c)
def box(d,xy,fill='white',outline=LINE,r=20): d.rounded_rectangle(xy,radius=r,fill=fill,outline=outline,width=2)
def wrap(s,n=32):
 out=[]; row=''
 for ch in s:
  if ch=='\n': out.append(row); row=''; continue
  if len(row)>=n: out.append(row); row=''
  row+=ch
 if row: out.append(row)
 return out
def lines(d,xy,s,n=24,width=30,c=NAVY,gap=12):
 x,y=xy
 for line in wrap(s,width): text(d,(x,y),line,n,c); y+=n+gap
 return y
def arrow(d,a,b,c='#7c91a2'):
 d.line([a,b],fill=c,width=4); ang=math.atan2(b[1]-a[1],b[0]-a[0]); z=12
 d.polygon([b,(b[0]-z*math.cos(ang-.55),b[1]-z*math.sin(ang-.55)),(b[0]-z*math.cos(ang+.55),b[1]-z*math.sin(ang+.55))],fill=c)
def diagram(name,title,rows):
 im=Image.new('RGB',(1500,180+len(rows)*150),'white'); d=ImageDraw.Draw(im); text(d,(45,30),title,36,b=True)
 for i,row in enumerate(rows):
  y=110+i*150; w=(1410-(len(row)-1)*25)//len(row)
  for j,(head,body) in enumerate(row):
   x=45+j*(w+25); box(d,(x,y,x+w,y+110),BG); text(d,(x+20,y+15),head,28,TEAL,True); lines(d,(x+20,y+57),body,22,max(8,int((w-35)/22)),gap=5)
  if i<len(rows)-1:
   next_w=(1410-(len(rows[i+1])-1)*25)//len(rows[i+1])
   parents=[45+j*(w+25)+w//2 for j in range(len(row))]
   children=[45+j*(next_w+25)+next_w//2 for j in range(len(rows[i+1]))]
   rail=y+130
   for cx in parents: d.line([(cx,y+110),(cx,rail)],fill='#7c91a2',width=4)
   d.line([(min(parents+children),rail),(max(parents+children),rail)],fill='#7c91a2',width=4)
   for cx in children: arrow(d,(cx,rail),(cx,y+150))
 im.save(ASSET/f'{name}.png')
 return name

diagram('functional','功能结构：双模式共用数据，训练策略分别配置',[
 [('两仪·TwoPoles','个人用户入口：登录、建档、模式选择')],
 [('大众模式','习惯目标 · 条件替换 · 基础跟练'),('运动员模式','跑步周期 · 计划对照 · 负荷记录')],
 [('训练执行','计数、计时、反馈、调整确认'),('饮食辅助','食物确认、类别分析、候选搭配'),('个人复盘','周报、趋势、数据导出')]
])
# Explicit client/server boundary; raw frames never flow into the server.
im=Image.new('RGB',(1500,780),'white'); d=ImageDraw.Draw(im)
text(d,(45,25),'系统总体架构与数据流向',36,b=True)
box(d,(45,90,1455,290),BG); text(d,(65,104),'浏览器客户端',27,TEAL,True)
box(d,(70,153,720,269)); text(d,(90,165),'Vue 界面与本地记录',27,b=True); text(d,(90,211),'双模式首页 · 训练日历 · 餐盘 · 周报',22)
box(d,(865,153,1430,269)); text(d,(885,165),'本地视觉处理',27,b=True); text(d,(885,207),'摄像头 → 关键点 → 动作状态与次数',21); text(d,(885,233),'原始画面默认不上传',18,MUTED)
arrow(d,(865,198),(725,198));text(d,(735,163),'计数摘要',20,MUTED)
text(d,(450,303),'用户确认后提交摘要 / HTTPS · JSON',23,TEAL)
box(d,(45,351,1060,760),BG); text(d,(660,359),'FastAPI 业务服务端',23,TEAL,True)
box(d,(70,390,1035,470));text(d,(90,410),'统一接口：身份校验 · 归属检查 · 请求去重',27,b=True)
arrow(d,(390,269),(390,390));arrow(d,(340,470),(340,520));arrow(d,(810,470),(810,520))
box(d,(70,520,540,611)); text(d,(90,535),'训练与规则服务',27,TEAL,True);text(d,(90,575),'条件过滤 · 计划版本 · 负荷统计',22)
box(d,(565,520,1035,611)); text(d,(585,535),'饮食与报告服务',27,TEAL,True);text(d,(585,575),'结构规则 · 缺失标记 · 周报聚合',22)
arrow(d,(340,611),(340,655));arrow(d,(810,611),(810,655))
box(d,(70,655,1035,740));text(d,(90,668),'SQLite 数据库',27,TEAL,True);text(d,(90,708),'账号 · 计划 · 训练 · 餐次 · 规则版本；事务写入',22)
box(d,(1100,520,1455,740));text(d,(1120,538),'静态资源',28,TEAL,True);lines(d,(1120,587),'固定版本姿态模型\n动作示范与食物目录\n随部署预置，按需加载',22,14,gap=14)
arrow(d,(1280,520),(1280,290));text(d,(1300,365),'本地加载',22,MUTED)
im.save(ASSET/'architecture.png')
diagram('flow','训练与饮食记录驱动下一次计划调整',[
 [('档案与约束','模式、目标、时段、器材；运动员另填专项与周期')],
 [('生成候选计划','规则过滤 → 安排训练 → 展示理由 → 用户确认')],
 [('训练记录','动作摘要、完成量、实际时长、主观用力'),('饮食记录','确认食物、份量或类别、餐次完整程度')],
 [('分析与预览','训练计划差异、负荷趋势；饮食结构与记录缺口')],
 [('用户决定','接受调整生成新版本；拒绝则保留原计划 → 下次训练')]
])
# ER schematic with explicit cardinality labels
im=Image.new('RGB',(1500,920),'white'); d=ImageDraw.Draw(im); text(d,(40,25),'主要数据关系示意',36,b=True)
nodes={'u':(590,110,'用户 User','账号与当前模式'),'p':(70,290,'档案 Profile','大众档案 / 运动员扩展'),'plan':(590,290,'计划 Plan','模式快照与版本'),'meal':(1120,290,'餐次 Meal','日期与记录完整度'),'item':(590,500,'计划项 PlanItem','动作或跑步安排'),'mi':(1120,500,'餐次明细 MealItem','食物与份量'),'s':(70,710,'训练 Session','实际训练与反馈'),'e':(590,710,'动作 Exercise','动作属性与识别配置'),'f':(1120,710,'食物 Food','类别及可用营养数据')}
for a,b,label in [('u','p','1 : 1'),('u','plan','1 : N'),('u','meal','1 : N'),('plan','item','1 : N'),('meal','mi','1 : N'),('item','e','N : 0..1'),('mi','f','N : 1'),('u','s','1 : N')]:
 x,y,_,_=nodes[a]; xx,yy,_,_=nodes[b]
 if a=='u' and b=='s':
  d.line([(590,155),(35,155),(35,760)],fill='#7c91a2',width=4); arrow(d,(35,760),(70,760)); text(d,(40,600),'1 : N',21,MUTED)
 else:
  arrow(d,(x+140,y+95),(xx+140,yy)); text(d,((x+xx)//2+145,(y+yy)//2+45),label,21,MUTED)
for x,y,h,b in nodes.values(): box(d,(x,y,x+280,y+100),BG); text(d,(x+14,y+15),h,25,TEAL,True); text(d,(x+14,y+58),b,19)
text(d,(55,855),'补充：计划项 1 : N 训练记录（临时训练可无计划项）；规则版本 1 : N 计划版本。',23,MUTED)
im.save(ASSET/'er.png')

def ui_base(title,mode):
 im=Image.new('RGB',(1500,940),BG); d=ImageDraw.Draw(im)
 d.rectangle((0,0,1500,90),fill=NAVY); text(d,(35,22),'两仪·TwoPoles',32,'white',True); text(d,(350,31),'个性化运动与饮食协同助手',24,'#c0d4e0')
 box(d,(1180,19,1460,70),'#27485c','#27485c',15); text(d,(1203,29),mode,24,'white'); d.polygon([(1423,40),(1439,40),(1431,49)],fill='white')
 d.rectangle((0,90,230,940),fill='white')
 for i,label in enumerate(['今日概览','训练计划','摄像头跟练','饮食记录','我的周报','个人设置']):
  yy=150+i*95
  if title.startswith(label): box(d,(15,yy-12,215,yy+55),'#e4f3f0','#e4f3f0',12)
  text(d,(40,yy),label,25,TEAL if title.startswith(label) else MUTED)
 text(d,(270,120),title,35,b=True); text(d,(270,174),'界面原型 · 示例数据',21,MUTED)
 return im,d

def button(d,xy,label,primary=True):
 box(d,xy,TEAL if primary else 'white',TEAL,12); text(d,(xy[0]+22,xy[1]+14),label,24,'white' if primary else TEAL,True)

im,d=ui_base('今日概览  /  从一小步开始','大众模式')
for x,val,label in [(270,'2 / 3','本周训练完成'),(675,'20 分钟','今日可用时间'),(1080,'基础体能','当前目标')]:
 box(d,(x,225,x+370,355)); text(d,(x+25,244),label,22,MUTED); text(d,(x+25,287),val,37,TEAL,True)
box(d,(270,385,1000,845)); text(d,(300,410),'今天的训练',30,b=True); text(d,(300,461),'宿舍 · 无器材 · 无跳跃',24,MUTED)
for i,(a,b) in enumerate([('动态热身','4 分钟'),('坐站训练','2 组 × 8 次'),('站姿弯举','2 组 × 10 次'),('舒缓整理','4 分钟')]):
 yy=522+i*55; text(d,(310,yy),a,25); text(d,(770,yy),b,23,MUTED)
button(d,(300,758,570,821),'开始跟练'); button(d,(600,758,960,821),'时间 / 器材有变化',False)
box(d,(1030,385,1450,650)); text(d,(1055,411),'饮食记录',29,b=True); lines(d,(1055,466),'早餐、午餐已记录。\n晚餐待补，全天分析未完成。',23,15,gap=7); button(d,(1055,562,1425,625),'记录这一餐',False)
box(d,(1030,680,1450,845)); text(d,(1055,702),'计划为什么这样安排',25,b=True); lines(d,(1055,748),'依据：20 分钟、无器材、入门经验。可随时调整。',23,15)
im.save(ASSET/'ui-public.png')

im,d=ui_base('今日概览  /  跑步专项训练台','运动员模式')
for x,val,label in [(270,'基础准备期','当前周期'),(675,'540 AU','近 7 日已记录负荷'),(1080,'3 / 5','本周训练记录')]:
 box(d,(x,225,x+370,355)); text(d,(x+25,244),label,22,MUTED); text(d,(x+25,288),val,35,TEAL,True)
box(d,(270,385,980,730)); text(d,(300,410),'周计划与实际完成',29,b=True)
text(d,(302,456),'时长 / 分钟',20,MUTED)
for tick in [0,30,60]:
 yy=638-tick*2.4;d.line((358,yy,942,yy),fill=LINE,width=2);text(d,(311,yy-13),str(tick),19,MUTED)
planned=[30,45,0,35,25,50,0];actual=[28,40,0,30,None,None,0]
for i,day in enumerate(['一','二','三','四','五','六','日']):
 x=372+i*81;text(d,(x+9,650),day,21,MUTED)
 if planned[i]:d.rounded_rectangle((x,638-planned[i]*2.4,x+24,638),5,fill=TEAL)
 if actual[i]:d.rounded_rectangle((x+29,638-actual[i]*2.4,x+53,638),5,fill='#93b3cd')
 if actual[i] is None:text(d,(x+29,614),'待',18,MUTED)
 if not planned[i]:text(d,(x+8,614),'休',18,MUTED)
text(d,(300,690),'绿色：计划   蓝色：实际   待：未记录   休：休息日',19,MUTED)
box(d,(1010,385,1450,720)); text(d,(1040,410),'今日训练',29,b=True); lines(d,(1040,465),'跑步 · 轻松训练\n计划 25 分钟\n强度目标：沿用个人计划\n结束后记录主观用力程度',23,16)
button(d,(1040,625,1418,689),'记录实际训练')
box(d,(270,750,1450,900)); text(d,(300,771),'有 1 条调整建议等待确认',27,b=True); text(d,(300,830),'周六时段缩短；锁定训练不变，确认后才保存。',23,MUTED); button(d,(1180,780,1415,855),'查看预览',False)
im.save(ASSET/'ui-athlete.png')

im,d=ui_base('摄像头跟练  /  坐站训练','大众模式')
box(d,(270,225,1040,825),'#e8eef3'); text(d,(295,245),'侧面入镜 · 请保证全身可见',23,MUTED)
# deliberate schematic stick figure, no fabricated camera capture
pts={'head':(650,360),'neck':(643,410),'hip':(610,565),'knee':(733,620),'ankle':(705,750),'shoulder':(643,435),'elbow':(720,477),'wrist':(777,440)}
for a,b in [('head','neck'),('neck','hip'),('hip','knee'),('knee','ankle'),('shoulder','elbow'),('elbow','wrist')]: d.line([pts[a],pts[b]],fill=TEAL,width=13)
d.ellipse((620,322,680,382),fill='#b4d6d1',outline=TEAL,width=5)
for p in pts.values(): d.ellipse((p[0]-8,p[1]-8,p[0]+8,p[1]+8),fill='white',outline=TEAL,width=3)
d.line((480,620,625,620,625,765),fill='#a1aeb9',width=10); text(d,(295,777),'关键点示意图；摄像头原始画面默认不上传',22,MUTED)
box(d,(1070,225,1450,530)); text(d,(1100,250),'第 1 组 / 共 2 组',26,b=True); text(d,(1100,315),'06 / 08',62,TEAL,True); text(d,(1100,409),'当前状态：起身阶段',23); text(d,(1100,455),'反馈：保持全身入镜',23,MUTED)
box(d,(1070,555,1450,825)); text(d,(1100,581),'完成后记录',27,b=True); lines(d,(1100,635),'完成次数可以修正。\n看不清时暂停计数，\n不输出姿势安全结论。',23,13); button(d,(1100,740,1415,802),'结束并反馈')
button(d,(270,850,495,913),'暂停',False); button(d,(530,850,835,913),'切换手动记录',False)
im.save(ASSET/'ui-pose.png')

im,d=ui_base('饮食记录  /  可编辑的餐盘','运动员模式')
box(d,(270,225,850,760)); text(d,(300,250),'午餐 · 先确认，再分析',29,b=True)
for i,(a,b) in enumerate([('米饭','谷薯类 · 约 1 碗'),('鸡蛋','蛋类 · 1 个'),('西兰花','蔬菜类 · 1 份'),('牛奶','奶类 · 250 mL')]):
 yy=325+i*83; box(d,(300,yy,815,yy+64),BG,BG,10); text(d,(320,yy+16),a,24,b=True); text(d,(450,yy+18),b,21,MUTED)
button(d,(300,678,815,738),'添加 / 修改食物',False)
box(d,(880,225,1450,550)); text(d,(910,250),'本餐结构',29,b=True)
for i,(a,col) in enumerate([('谷薯类',TEAL),('蔬菜类','#5a9b6c'),('蛋类','#d7a54b'),('奶类',BLUE)]):
 x=912+(i%2)*245; y=320+(i//2)*65; d.ellipse((x,y,x+20,y+20),fill=col); text(d,(x+34,y-5),a,24)
lines(d,(910,464),'已确认 4 类食物。份量为用户估计，暂不提供精确营养结论。',22,22)
box(d,(880,580,1450,760)); text(d,(910,605),'训练日记录',27,b=True); text(d,(910,657),'今日训练 17:00 · 午餐 12:10',23); text(d,(910,701),'仅用于回顾，不判定恢复是否充分。',23)
box(d,(270,790,1450,904)); text(d,(300,810),'全天分析待完善',26,b=True); text(d,(300,858),'早餐、晚餐尚未确认完整。缺少记录不会被当成“没有吃”。',23,MUTED); button(d,(1200,813,1420,877),'保存餐次')
im.save(ASSET/'ui-food.png')

# Template contract and package inventory.
with ZipFile(REF) as z: inventory={n:hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist()}
(ROOT/'qa'/'template-inventory.json').write_text(json.dumps(inventory,indent=2),encoding='utf-8')
(ROOT/'artifact.md').write_text(f'''# Template execution contract
Reference: {REF}\nSHA256: {hashlib.sha256(REF.read_bytes()).hexdigest()}\nReference: 10 pages, one A4 portrait section. Render: qa/template.pdf; qa/template-render.
Page geometry: 11906 x 16838 twips; margins top1191 right1247 bottom1134 left1361; header/footer454; preserve sectPr and all header/footer parts.
Styles: preserve styles.xml byte for byte. Course Heading 1 16pt blue1F4E79; Heading2 13pt blue2F75B5; Heading3 source-defined. Normal DengXian10.5pt line1.35 after4pt. Cover original four paragraph patterns, original title and visual rule. Source styling overrides generic defaults.
Slots: cover p0..p3 metadata; all numbered headings retained; instructional body and example table rows replaced with project content; generic XXX modules expanded to eight modules, each with four prescribed subheadings. Source heading and normal components cloned. Additional figures placed inline, labels in source Normal. Table components cloned, widths expanded to content width and rows repeat across pages. All instructional text intentionally removed.
Figures: four diagrams plus four UI wireframes, marked as prototype/sample data. No software implementation represented as completed.
Preserve-only: all original package parts except document.xml, document.xml.rels, [Content_Types].xml; add PNG media only. No original media removed. Keep numbered chapter order1..9 then references. Cover member names/student IDs intentionally blank. All contents fully populated; performance numbers are proposed acceptance targets.
Intentional changes: content expansion and pagination, table widths, repeated header/row keep settings; cover metadata indent reduced for project name; content graphs and UI images added as requested. Original files unchanged.
''',encoding='utf-8')

doc=Document(REF); original_paras=list(doc.paragraphs); table_template=deepcopy(doc.tables[2]._tbl)
normal_template=deepcopy(original_paras[6]._p)
heading_templates={1:deepcopy(original_paras[4]._p),2:deepcopy(original_paras[5]._p),3:deepcopy(original_paras[96]._p)}
cover=[deepcopy(p._p) for p in original_paras[:4]]
for el in list(doc._element.body):
 if el.tag!=qn('w:sectPr'): doc._element.body.remove(el)
def insert(el): doc._element.body.insert(len(doc._element.body)-1,el)
def setp(el,s):
 for child in list(el):
  if child.tag!=qn('w:pPr'): el.remove(child)
 r=OxmlElement('w:r'); t=OxmlElement('w:t'); t.text=s; t.set(qn('xml:space'),'preserve'); r.append(t); el.append(r)
 return el
for i,p in enumerate(cover):
 if i==2: setp(p,'程序设计综合实验 · 双模式运动与饮食辅助系统')
 if i==3:
  ppr=p.find(qn('w:pPr')); ind=ppr.find(qn('w:ind'))
  if ind is not None: ind.set(qn('w:left'),'1200'); ind.set(qn('w:right'),'650')
  setp(p,'')
  rr=p.find(qn('w:r')); p.remove(rr)
  for j,line in enumerate(['项目名称：两仪·TwoPoles','（个性化运动与饮食协同助手）','组员1：________（学号：________）','组员2：________（学号：________）','组员3：________（学号：________）','完成日期：2026年10月13日']):
   r=OxmlElement('w:r'); rp=OxmlElement('w:rPr'); sz=OxmlElement('w:sz');sz.set(qn('w:val'),'22');rp.append(sz);r.append(rp);t=OxmlElement('w:t');t.text=line;r.append(t)
   if j<5: r.append(OxmlElement('w:br'))
   p.append(r)
  br=OxmlElement('w:br');br.set(qn('w:type'),'page');r.append(br)
 insert(p)

def para(s,keep=False):
 el=setp(deepcopy(normal_template),s); insert(el)
 from docx.text.paragraph import Paragraph
 p=Paragraph(el,doc._body); p.paragraph_format.keep_with_next=keep; p.paragraph_format.widow_control=True; p.paragraph_format.keep_together=True
 return p
def h(s,level=2,page=False):
 if page:
  prev=doc._element.body[-2]
  if prev.tag==qn('w:p') and not ''.join(prev.itertext()).strip() and not prev.findall('.//'+qn('w:drawing')): doc._element.body.remove(prev)
 el=setp(deepcopy(heading_templates[level]),s); insert(el)
 from docx.text.paragraph import Paragraph
 p=Paragraph(el,doc._body)
 if page: p.paragraph_format.page_break_before=True
 return p
def table(headers,rows,widths=None):
 t=doc.add_table(rows=1,cols=len(headers)); t._tbl.remove(t._tbl.tblPr); t._tbl.insert(0,deepcopy(table_template.tblPr)); t.autofit=False
 total=16.4
 if widths is None: widths=[total/len(headers)]*len(headers)
 for col,w in zip(t.columns,widths): col.width=Cm(w)
 for cells,data in [(t.rows[0].cells,headers)]+[(t.add_row().cells,row) for row in rows]:
  for c,s,w in zip(cells,data,widths):
   c.width=Cm(w); c.text=str(s)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(3); p.paragraph_format.space_before=Pt(3); p.paragraph_format.line_spacing=1.15
    for r in p.runs: r.font.size=Pt(9.5)
 for idx,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr(); pr.append(OxmlElement('w:cantSplit'))
  if idx==0: pr.append(OxmlElement('w:tblHeader'))
  for c in row.cells:
   cp=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'1F4E79' if idx==0 else ('F3F7FA' if idx%2 else 'FFFFFF'));cp.append(sh)
   borders=OxmlElement('w:tcBorders')
   for side in ['top','bottom','left','right']:
    x=OxmlElement('w:'+side);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
   cp.append(borders)
   if idx==0:
    from docx.shared import RGBColor
    for p in c.paragraphs:
     for r in p.runs: r.font.bold=True;r.font.color.rgb=RGBColor(255,255,255)
 para('')
 return t

def fig(name,caption,width=16.1):
 p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
 p.add_run().add_picture(str(ASSET/(name+'.png')),width=Cm(width))
 p=para(caption);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 for r in p.runs:r.font.size=Pt(9)

h('1. 项目概述',1)
h('1.1 项目背景与简介')
para('大众健身用户常面临训练动作选择困难、计划与实际时间不符、饮食记录难以转化为具体反馈等问题；具有训练基础的运动员则需要围绕专项和周期整理训练安排、比较计划与实际完成情况，并持续观察个人训练记录。两类用户的目标与信息需求不同，统一的功能入口和反馈方式难以同时满足其使用习惯。本项目拟开发“两仪·TwoPoles（个性化运动与饮食协同助手）”，面向成年普通健身者及校队、业余竞技运动员。大众模式提供条件化训练推荐、摄像头跟练和饮食结构提示；运动员模式提供周期计划、跑步专项记录及负荷趋势。系统通过“建档—计划—执行—反馈—调整”形成完整使用流程，采用可解释规则和用户确认机制，帮助用户记录、理解并调整日常安排。首期不提供医疗诊断、伤病康复处方或专业竞技决策，饮食分析用于一般性结构参考。')
h('1.2 项目目标')
para('项目总体目标是实现可独立演示、可重复验证的双模式个人运动辅助系统。以下指标为计划验收目标，不代表已经完成开发或获得实测结果。')
table(['目标','可检查的预期结果'],[
('双模式服务','同一账号可进入大众或运动员模式；切换前展示变化并确认，历史记录及原计划版本不丢失。'),
('完整训练流程','支持建立档案、生成或录入计划、执行、反馈、调整预览与确认；每次计划修改保留原因和版本。'),
('可交互的视觉功能','支持侧面坐站及站姿弯举两个演示动作；在固定机位、单人、无遮挡测试集上，至少90%的测试片段计数误差不超过1次。'),
('饮食结构辅助','支持手工录入并确认食物类别；缺失记录不当作未摄入；每项结构提示可查看依据和适用范围。'),
('运动员专项功能','首期覆盖跑步计划与实际记录，统计时长、距离、主观用力和个人负荷趋势；不从负荷值推断伤病概率。'),
('可验证的交付','提交源代码、部署说明、测试集与测试报告、文档和演示视频；核心流程不依赖付费模型API。')],[3.1,13.3])
h('1.3 用户与使用场景')
table(['用户类型','使用场景','主要需求'],[
('大众用户','宿舍、家庭或健身房进行一般体能训练','低门槛建档；按时间和器材替换动作；计数、习惯记录、饮食类别反馈。'),
('运动员用户','成年校队或业余竞技运动员记录个人训练，首期以跑步为示范','周期目标、锁定既有训练、计划与实际对照、个人负荷与训练日用餐时间线。'),
('管理员','维护经审核的动作库、食物目录与分析规则','资料版本、来源维护、异常日志；默认无权查看用户私密饮食和训练详情。')],[2.6,6.0,7.8])
para('模式是同一用户的使用配置，不是不同权限等级，也不通过体重、成绩或收费自动划分。首期仅面向18岁及以上用户；已有伤病、孕期或有特殊膳食管理需求者可保留记录功能，但不进入自动训练推荐流程。教练多人管理端不属于首期范围。')
table(['比较维度','大众模式','运动员模式'],[
('目标与输入','建立习惯、一般体能；经验、时段、器材、偏好','专项、训练经验、周期阶段、目标日期、既有计划；可录入体重用于个人趋势'),
('计划方式','从基础模板生成一周安排，支持动作替换','录入或复制个人周期计划，首期支持4周跑步安排；保留锁定训练'),
('执行与反馈','跟练计数、完成情况、训练后感受','跑步距离/时长、实际与计划差异、session-RPE记录；基础动作跟练可共用'),
('饮食呈现','本餐/当天食物类别与记录完整程度','增加训练日和休息日标记、用餐与训练时间线；不自动开具竞技营养处方'),
('调整策略','根据时间、器材与反馈提供候选替换','只提出局部调整建议及影响，不擅自改变周期目标或教练既有安排')],[2.6,6.9,6.9])
h('1.4 系统运行环境')
para('客户端为Windows 10/11或macOS电脑上的现代浏览器，移动端支持基础记录和查看。摄像头计数以电脑端Chrome/Edge为首期验收环境；需要720p摄像头、单人完整入镜和稳定光照。基准测试设备拟采用四核CPU、8GB内存，不要求独立GPU；实际设备型号和浏览器版本在测试报告中记录。')
para('服务端使用Python运行环境和SQLite数据库，可部署在Windows或Linux。校园局域网或本机即可演示；公网部署使用HTTPS，摄像头依赖安全上下文与用户授权。本地演示使用localhost。模型和规则资源提前下载并固定版本，初次准备后核心演示不依赖外网。')
h('1.5 开发环境与开发工具')
table(['类别','工具/技术','版本或说明'],[
('前端','Vue 3、TypeScript、Vite','响应式网页；版本在立项时锁定并保留依赖清单'),('后端','Python 3.11、FastAPI','业务接口、输入校验和规则服务'),('视觉','MediaPipe Tasks Vision、Canvas','浏览器端姿态推理；模型和包固定版本及校验值'),('数据','SQLite 3、SQLAlchemy、Alembic','单机部署、结构迁移与事务；保留升级接口'),('图表与测试','ECharts、pytest、Playwright','趋势图、算法与接口测试、关键页面流程检查'),('开发协作','VS Code、Git','分支评审、任务记录和问题跟踪')],[2.2,6.2,8.0])
h('1.6 相关系统或技术调研')
para('本项目选择两类用户产品及一项基础技术进行调研，重点借鉴可实现的交互和记录方式，不以“首创健身平台”作为创新性依据。')
table(['系统或技术','可借鉴内容','本项目的定位与取舍'],[
('Garmin Connect训练负荷相关功能 [1]','训练记录与负荷趋势的组织方式','首期采用手工时长与主观用力记录；不复制依赖设备的专有算法，不宣称等同Garmin指标。'),('MyFitnessPal饮食日记 [2]','按餐次记录食物并回顾日常摄入','采用小型来源明确的食物目录，优先解释结构与缺失；不追求庞大的商业食物数据库。'),('MediaPipe Pose Landmarker [3]','人体关键点估计与浏览器视觉能力','在关键点上自行实现状态机计数、低置信度暂停和事件摘要；不把关键点检测当作姿势安全诊断。')],[4.0,5.2,7.2])
para('项目区分度在于双模式的可解释训练流程、低门槛跟练记录和训练饮食时间线的统一复盘。系统以个人运动过程为主线，饮食模块提供配套结构参考；将交互式训练执行、条件替换、计划变更确认作为主要实现重点。')

h('2. 需求分析',1,True)
h('2.1 功能需求')
para('“必须”项组成首期验收范围；“可选”项不影响核心流程验收。运动员模式必须交付，不作为大众模式的后期扩展。')
features=[
('F01','账号与数据控制','注册、登录、退出；用户数据隔离；导出个人记录与删除账号。','必须'),
('F02','双模式档案','填写共同档案及运动员扩展信息；切换模式先确认，保留历史模式快照。','必须'),
('F03','训练计划','大众一周模板推荐；运动员4周跑步计划录入/复制；日历查看、锁定和版本保存。','必须'),
('F04','条件替换','按时长、场地、器材过滤候选动作，解释替换原因；无候选时如实提示。','必须'),
('F05','摄像头跟练','两个指定动作的关键点可视化、计数与阶段提示；支持暂停和手工修正。','必须'),
('F06','训练反馈','记录实际组次/时长、完成原因和主观用力；运动员增加距离与训练类型。','必须'),
('F07','计划调整','基于反馈生成局部调整预览；锁定项保持不变；确认后生成新版本。','必须'),
('F08','负荷与趋势','运动员按session-RPE与实际分钟数计算记录负荷；展示周汇总及缺失状态。','必须'),
('F09','饮食记录','按餐次选择食物、类别、可选份量；编辑和确认完整程度，不强制拍照。','必须'),
('F10','结构分析','给出食物类别覆盖、记录缺口和替换候选；运动员增加训练日时间线。','必须'),
('F11','个人周报','展示计划完成、训练趋势、饮食记录完整度及依据；可导出数据。','必须'),
('F12','资料与规则管理','维护至少20个基础动作、40个食物条目及规则来源和版本；停用不适用条目。','必须'),
('F13','照片辅助录入','识别候选食物并由用户确认；无法确认则回退手工录入。','可选'),
('F14','食堂配餐实验室','对自建菜单按预算、偏好和食物类别生成可比较的搭配。','可选'),
('F15','知识问答','检索已审核知识条目并显示来源；不得直接改写计划或输出诊疗结论。','可选'),
('F16','视频回放与设备接入','用户单独授权保存本地片段；探索CSV训练导入，穿戴设备直连暂不承诺。','可选')]
table(['编号','功能名称','功能描述','优先级'],features,[1.15,2.9,10.95,1.4])
h('2.2 功能结构或用例分析')
fig('functional','图2-1 系统功能结构图')
para('典型用例一：大众用户选择“宿舍、无器材、20分钟”，确认基础计划，进入跟练并保存完成次数；系统依据反馈生成下次安排预览。典型用例二：运动员录入跑步周期与锁定训练，记录距离、时长及主观用力，查看周趋势，确认或拒绝局部调整。典型用例三：用户确认餐次中的食物，查看类别提示；记录不足时仅分析已确认部分。')
para('管理员仅维护内容与规则。切换模式、保存调整、删除记录和上传可选媒体等操作均需要明确操作入口；关键动作给出成功或失败反馈。')
h('2.3 输入与输出需求')
table(['功能','输入','输出'],[
('建档与切换','年龄确认、模式、目标、时段、器材；专项与周期','档案校验结果、模式变化说明及对应首页'),('计划与替换','档案、动作库、既有计划、锁定项、可用时间','候选计划、无法满足的条件、变更对照和版本'),('视觉跟练','已授权摄像头视频、动作类型、机位校准','关键点、动作阶段、计数、置信度提示、事件摘要'),('运动员记录','实际分钟数、可选距离、0～10主观用力、训练日期','单次负荷、个人周趋势、计划实际对照'),('饮食分析','确认食物类别、可选份量、餐次完整度与时间','类别分布、分析依据、缺口说明和可选替换'),('个人报告','日期范围、训练与饮食记录','页面图表、统计口径、个人CSV/JSON导出')],[2.6,7.0,6.8])
h('2.4 非功能需求')
para('（1）易用性：首次使用分步建档；两种模式使用相同导航；颜色之外同时使用文字表达状态。主要按钮支持键盘操作；摄像头拒绝授权时仍能手工记录。五名目标用户完成“建档—记录训练—查看分析”任务的成功人数目标不少于四人，并记录耗时与问题。')
para('（2）性能：在1.4节基准设备及本地网络上，以30次重复测试统计，普通列表接口P95响应不超过1秒，单次计划生成或调整P95不超过3秒；测试数据为100个账号和累计1万条记录，10个并发请求。姿态功能以640×480输入进行性能测试，目标有效推理频率不低于10次/秒；达不到时降采样或提示手工记录，不伪造实时结果。')
para('（3）可靠性：输入错误、网络中断及模型加载失败不能导致已保存数据损坏。写入使用事务与请求去重；计划修改采用版本检查。数据库每日备份，验收至少完成一次恢复演练；空数据与缺失数据均有明确显示。')
para('（4）安全性：密码采用成熟库的加盐哈希存储；所有个人数据接口校验记录归属。摄像头原始视频默认只在浏览器处理且不上传；只保存用户确认的训练摘要。日志不得包含密码、令牌和完整饮食明细。删除账号应同步删除关联记录；离线备份最长30天轮换清理并说明保留期限。')
para('（5）其他要求：分析可追溯到规则版本与来源。饮食记录不足时不生成全天营养结论；疼痛或明显不适触发停止自动加量提示。运动员负荷仅用于个人记录比较，不输出伤病预测、恢复评分或比赛成绩保证。以上限制贯穿两种模式及可选问答功能。')

h('3. 系统概要设计',1,True)
h('3.1 系统总体架构')
fig('architecture','图3-1 系统总体架构图')
para('采用浏览器/服务端架构。前端负责表单、日历、图表及本地姿态计算；FastAPI服务端负责身份校验、计划版本、规则分析和数据持久化。视觉原始画面不经过业务接口，前端仅在用户确认后提交计数和训练摘要。业务层通过统一的数据访问模块读写SQLite，规则及静态资源采用版本化管理。[3][7][8]')
h('3.2 功能模块划分')
modules=[('M1','用户档案与双模式模块','账号、建档、模式切换、个人数据控制；F01～F02'),('M2','训练计划与调整模块','大众模板、运动员周期、替换、预览、确认与锁定；F03～F04、F07'),('M3','姿态跟练模块','授权、机位校准、关键点、两个动作计数与异常降级；F05'),('M4','训练记录与负荷模块','完成记录、主观用力、距离时长、个人负荷汇总；F06、F08'),('M5','饮食记录与结构分析模块','餐次确认、类别分析、来源说明、训练日时间线；F09～F10'),('M6','报告与可视化模块','计划实际对照、个人趋势、记录完整度与导出；F11'),('M7','内容规则与系统支撑模块','资料审核、版本、权限、日志、备份及恢复；F12'),('M8','可选交互扩展模块','照片辅助、配餐实验室、问答和可选回放；F13～F16')]
table(['模块编号','模块名称','主要功能'],modules,[1.6,5.1,9.7])
h('3.3 模块关系及数据流程')
fig('flow','图3-2 主要业务数据流程图')
para('M1输出档案与当前模式给M2、M5；M2引用M7的动作与规则，生成候选计划，用户确认后保存。M3产生动作摘要，交给M4与手工反馈合并；M4保存实际记录，供M2提出调整、M6生成趋势。M5保存确认餐次与分析依据，M6将其与训练日期关联。M8只提供可选候选或展示，不绕过M1权限与M2/M5的确认环节。')
para('模式切换只更新后续操作的界面配置；每条计划保留创建时的模式与规则版本。未完成旧计划继续可见，只有用户明确复制或调整后才形成新模式下的计划，防止静默覆盖。')
h('3.4 主要技术方案')
h('3.4.1 图形界面技术',3)
para('采用Vue 3和TypeScript构建响应式网页，使用日历和训练卡展示计划，Canvas叠加关键点，ECharts绘制趋势。训练页面把实时计数与数据保存分离，避免网络波动打断本地跟练；重操作可放入Web Worker，手机端若性能不足则保留记录与报告功能。[3][7]')
h('3.4.2 数据存储技术',3)
para('SQLite适合课程项目单实例部署，减少环境配置；使用SQLAlchemy统一数据访问并通过迁移脚本管理结构。训练、饮食和计划版本采用关系表，分析解释保留结构化摘要和规则版本。大体积视频不写入数据库。计划确认在同一事务中校验旧版本并写入新版本；数据库备份采用可保持一致性的备份方式，不直接复制正在写入的文件。')
h('3.4.3 核心算法或处理方法',3)
para('（1）条件化训练推荐：先按模式、经验、器材、场地、时间及用户限制过滤动作，再依据训练模板组合候选。排序考虑用户偏好、最近完成情况及相对原计划的改动量；约束无法满足时返回具体冲突，不强行安排。运动员首期以用户录入的跑步周期为基准，仅提出局部时间或辅助动作调整，不自动制定专业训练强度。')
para('（2）视觉计数：以MediaPipe关键点为输入，做置信度过滤和平滑，选取与动作相关的角度或位移，通过“起始—运动—返回”的状态机计数。采用迟滞区间和最短阶段持续时间减少抖动误计。阈值基于标注样本与机位校准确定，不能当作通用标准姿势；保存模型版本、动作配置和人工修正次数。[3]')
para('（3）运动员负荷记录：采用session-RPE方法，将训练后填写的整体主观用力程度（0～10）与实际训练时长（分钟）相乘，得到任意单位AU的单次记录值；同一统计周期相加形成周汇总。该方法参考相关研究，产品应统一填写提示与口径，并仅用于同一用户的趋势回顾。未填写强度的训练标为“负荷缺失”，不能按0计算或据此判定伤病风险。[6]')
para('（4）饮食结构规则：先区分“类别确认”“份量估计”和“记录完整度”，再对已确认食物进行类别覆盖分析。一般建议参考《中国居民膳食指南（2022）》；对于运动员，增加训练时间线与用户自定目标对照，不套用未经审核的竞技营养阈值。混合菜可人工拆分，无法判断的类别保持未知；只有具有明确来源与单位的食物数据才参加营养估算。[4]')
para('（5）验证方法：准备至少40段经同意采集或授权的动作测试片段，两个动作各不少于20段，含完整动作、半程、停顿和遮挡；至少5名参与者，调参样本与验收样本按参与者隔离。人工复核计数作为基准；另设20组计划约束案例和20组饮食/缺失案例。报告误计、漏计、失败条件与运行环境，不只展示成功片段。')

h('4. 模块概要设计',1,True)
para('各模块按模板规定的“模块功能、主要输入、主要输出、基本处理流程”描述；本阶段不展开函数、类及源代码。')
module_details=[
('用户档案与双模式模块','负责账号、成年范围确认、共同档案和运动员扩展档案；提供模式切换、导出与删除入口。模式只影响业务呈现，不改变数据所有权。','账号凭据、目标、经验、可用时段、器材、用户限制、专项、周期阶段及模式切换确认。','已校验的个人档案、当前模式、可推荐状态、变更说明；导出文件或删除结果。','注册/登录 → 校验输入与范围 → 保存档案 → 展示模式差异 → 用户确认 → 加载对应首页。切换时检查未完成计划，保留历史并提示后续处理。'),
('训练计划与调整模块','负责基础模板推荐、4周跑步计划录入/复制、日历、锁定训练和版本化调整；规则无解时返回原因。','M1档案、可用时段、动作/模板库、既有计划、M4实际反馈及规则版本。','候选计划、动作替换列表、变更对照、解释记录和用户确认后的计划新版本。','读取约束 → 过滤候选 → 组织时段与训练项 → 检查冲突 → 展示计划。调整时先固定完成项和锁定项，再生成局部变动预览；确认时检查版本，一次事务保存。'),
('姿态跟练模块','负责授权和机位引导、两个动作的关键点显示、状态转换、计数、暂停与手工修正；不输出医学或姿势安全结论。','本地视频帧、选定动作、模型与动作规则版本、校准结果、开始/暂停指令。','实时阶段、次数、可靠性提示、动作事件摘要、用户确认后的最终次数。','检查设备权限 → 加载模型 → 指导完整入镜 → 推理和过滤关键点 → 平滑及状态机识别 → 更新计数 → 用户结束并校正 → 提交摘要。遮挡期间冻结判断，恢复后重新进入稳定起始状态。'),
('训练记录与负荷模块','统一保存跟练与手工训练；运动员记录跑步距离、时长、训练类型与session-RPE，计算个人可比的记录负荷。','计划项或临时训练标记、动作摘要、实际时间、组次/距离、主观用力、未完成原因及不适反馈。','持久化训练记录、完成比例、单次AU、汇总结果与缺失提示，向计划模块提供事实反馈。','读取或建立训练条目 → 合并手工与视觉摘要 → 校验单位和范围 → 用户确认 → 事务去重写入 → 计算统计。主观用力缺失时保留训练量但不填补负荷；不适反馈优先阻止自动加量。'),
('饮食记录与结构分析模块','负责餐次建档、食物类别与份量确认、结构分析、记录完整度和训练日时间线；必要时显示一般适用范围。','用户选定食物、混合菜拆分、可选份量、餐次时间与完整度、食物目录及规则来源；运动员训练日期。','确认后的餐次、类别覆盖、缺失项、规则解释和可选替代；具备来源和份量时才有估算营养值。','选择餐次 → 录入或修正食物 → 确认类别/份量 → 标记完整程度 → 引用版本规则分析 → 展示结构与未知项 → 保存。删除或修改食物后重新分析，不保留过期结论。'),
('报告与可视化模块','聚合计划和实际训练、负荷趋势与饮食记录完整度；提供同一用户不同周期对照和结构化导出。','日期范围、计划版本、训练记录、已确认餐次、缺失标记及用户导出请求。','个人周报、计划实际对照、可追溯图表和CSV/JSON文件；空数据时给出明确提示。','校验访问权限 → 按时区和日期选择数据 → 去重聚合 → 计算具有完整分母的指标 → 展示说明 → 可选导出。没有安排训练的一周不计算完成率；计划外训练单独统计。'),
('内容规则与系统支撑模块','维护动作、食物、规则来源、发布版本及停用状态；提供权限校验、脱敏日志、备份与恢复。','管理员审核的资料及来源、规则配置、备份计划和系统异常事件。','可追溯资源版本、运行日志、备份文件与恢复结果；为其他模块提供稳定基础服务。','录入资料 → 校验来源与字段 → 审核发布 → 新计划引用新版本，旧记录保留旧版本 → 监控异常 → 备份及恢复验证。管理员默认只见内容资料和脱敏统计。'),
('可选交互扩展模块','承载照片辅助录入、食堂组合推荐、带来源问答和用户自愿保存的回放；独立启停，不改变必须功能的可用性。','明确授权的照片/片段、菜单与预算、查询文本或规定格式的CSV；外部模型可用性状态。','待确认的候选条目、配餐比较、引用来源回答或可回放文件；失败时提供手工入口。','检查开关及授权 → 校验文件与请求 → 调用适配器 → 验证结果结构 → 显示不确定项 → 用户确认后进入M2/M5/M4。接口失败不生成假结果，不把媒体自动保存到服务器。')]
for i,(name,fun,inp,out,flow) in enumerate(module_details,1):
 h(f'4.{i} {name}',2,page=i in (3,5,7))
 for j,(label,content) in enumerate([('模块功能',fun),('主要输入',inp),('主要输出',out),('基本处理流程',flow)],1): h(f'4.{i}.{j} {label}',3);para(content)

h('5. 数据概要设计',1,True)
h('5.1 主要数据说明')
table(['数据名称','数据内容','来源','用途'],[
('用户与档案','账号标识、密码哈希、模式、目标、器材、专项、周期','用户输入','身份验证、个性化约束'),('动作与模板','动作属性、示范说明、识别配置、适用条件','审核资料','训练推荐和视觉计数'),('计划与版本','日期、计划项、锁定状态、模式和规则快照','规则生成或用户录入','执行、调整与变更追溯'),('训练记录','实际时长、次数、距离、主观用力、不适反馈','视觉摘要或手工确认','进度与个人负荷统计'),('食物与餐次','类别、可选营养来源、份量、时间、完整度','食物目录和用户确认','结构分析与时间线'),('分析与规则','解释文本、引用来源、规则版本、未知项','规则计算','复盘与结果核对'),('系统日志','匿名事件标识、错误类型、时间、版本','系统运行','故障定位；不记录私密正文')],[2.5,6.1,3.5,4.3])
h('5.2 数据库或数据结构设计')
para('采用关系数据库。下表说明主要实体用途和关系，字段约束及索引细节留待详细设计。所有个人业务表直接或经父表关联用户；删除、查询和导出均按用户边界执行。')
table(['主要数据表','主要用途','关系与约束'],[
('User / Profile','账号、共同档案及运动员扩展','User与Profile为1:1；模式保留变更时间'),('Exercise / Template','动作属性、基础训练模板','模板可引用多个动作；停用条目不可用于新推荐'),('Plan / PlanItem','版本化计划及具体训练安排','User 1:N Plan；Plan 1:N PlanItem；跑步项可不引用动作'),('TrainingSession / Feedback','一次实际训练及反馈','User 1:N Session；PlanItem 1:N Session；临时训练允许无PlanItem'),('Food / Meal / MealItem','食物目录、餐次和明细','User 1:N Meal；Meal 1:N MealItem；MealItem N:1 Food'),('RuleVersion / AnalysisRecord','规则发布版本及分析快照','每条分析关联对应输入版本和规则；输入修改使旧结果失效'),('AuditEvent','脱敏安全事件与资料变更','记录操作者与动作，不存密码、令牌和敏感正文')],[5.0,5.1,6.3])
fig('er','图5-1 主要数据实体关系图')
para('训练时长统一存分钟，距离统一存米，显示时可转换；日期同时保留用户时区，周报按用户本地周统计。份量未知使用空值，不使用0代替；模型计数和人工修正值分别保留。原始视频默认不持久化，报告由业务记录生成，避免多份统计数据失去一致性。')

h('6. 用户界面与外部接口概要设计',1,True)
h('6.1 用户界面原型')
para('以下为拟实现界面的原型图，数值均为交互示例，不代表真实用户或实测结果。页面导航保持一致，模式切换先展示差异并确认；运动员专属数据以文字标识，避免仅靠颜色区分。')
fig('ui-public','图6-1 大众模式首页原型')
para('大众首页突出“开始跟练”和“条件变化”两个入口，减少初学者选择负担。饮食卡片显示记录缺口；无训练计划时引导建档，未授权摄像头时可进入手工计时与记录。')
h('6.1.1 运动员模式首页',3,True)
fig('ui-athlete','图6-2 运动员模式训练台原型')
para('训练台展示周期、计划实际对照和个人负荷。缺少主观用力的数据标记为未计入；AU为记录指标，不显示“最佳负荷区间”或恢复诊断。调整建议点击后展开旧值、新值和原因，支持接受、拒绝及保留锁定项。')
h('6.1.2 摄像头跟练界面',3)
fig('ui-pose','图6-3 摄像头跟练原型')
para('界面同时呈现动作、组次和计数状态。首次进入展示机位示范和权限说明；未完整入镜或置信度不足时提示暂停计数。用户可手工修正、结束训练和填写反馈；图中骨架为示意绘制，不是实际识别结果。')
h('6.1.3 饮食结构分析界面',3,True)
fig('ui-food','图6-4 可编辑餐盘与训练日时间线原型',width=14.2)
para('先确认食物，再分析结构；份量和营养估算分别标注依据。大众模式显示类别与记录完整度，运动员模式增加训练时间对照。当前未记录的餐次不会被判定为未摄入；可选拍照识别只预填候选，不直接生成结论。')
h('6.2 外部接口')
table(['接口或资源','用途与主要数据','依赖及失败处理'],[
('浏览器摄像头接口','getUserMedia采集本地视频；需用户明确授权','HTTPS/localhost；拒绝授权则手工记录'),('MediaPipe模型资源 [3]','加载本地固定版本模型，输出人体关键点','随部署准备资源；加载失败关闭视觉功能'),('本系统业务接口','前端以JSON交换档案、计划、训练、餐次和报告','统一登录鉴权、数据校验与错误提示；重试不重复写入'),('审核后的本地数据文件','JSON/CSV初始化动作、食物目录、规则和来源','校验格式、单位及版本；异常条目拒绝导入'),('可选外部模型API','照片候选识别或带来源文本解释','默认关闭；单独授权并最小化传输，超时回退手工入口'),('导出与可选导入文件','个人CSV/JSON导出；规定结构的训练CSV导入','不承诺直连商业穿戴平台；导入先预览并去重')],[4.0,6.1,6.3])

h('7. 异常情况与处理方案',1,True)
table(['异常情况','预期处理方式'],[
('必填项为空、数值越界或单位错误','定位到字段并提示允许范围；保留其他已填写内容，不保存无效记录。'),('不符合自动推荐适用范围','显示一般记录入口；暂停自动训练建议，不将运动员模式作为绕过入口。'),('模式切换时存在未完成计划','展示旧计划保留方式，明确切换影响；未确认则取消切换。'),('时间、器材或锁定项导致计划无解','列出冲突条件，提供修改时段/选择候选的入口；不擅自解除锁定。'),('多页面同时修改同一计划','版本不一致则拒绝覆盖，提示刷新并重新比较变更。'),('摄像头拒绝、模型加载失败','说明原因并提供手工计时、次数录入；其他功能继续可用。'),('遮挡、多人入镜或关键点不稳定','暂停相关动作判断，提示重新入镜和校准；不累计可疑次数。'),('运动中反馈疼痛或明显不适','提示停止当前训练并寻求适当专业帮助；取消自动加量建议，保留事实记录。'),('主观用力、份量或餐次记录缺失','明确标记未知；负荷或营养汇总注明覆盖范围，不以0补齐。'),('重复提交、断网或接口超时','幂等写入、提示重试；训练结束前保留当前页摘要，提醒不要关闭页面。'),('数据库写入失败或磁盘不足','回滚事务，给出失败提示并记录脱敏错误；恢复后可重新提交。'),('账号越权、会话失效','拒绝访问且不返回记录内容；要求重新登录，记录安全事件。'),('可选识别、问答产生不可信结果','不自动保存，不修改计划；提示确认或切回手工/规则功能。'),('删除记录后报告仍显示旧结果','使关联分析缓存失效并重算；删除账号执行级联清理及备份轮换。')],[5.2,11.2])

h('8. 项目开发计划',1,True)
para('按三人小组、六周开发安排制定相对计划，第一周从课程确认立项后开始。组员1负责前端与交互，组员2负责后端、数据与部署，组员3负责视觉、规则与算法验证；成员共同进行资料审核和集成测试。具体姓名与学号在封面补充，不预设实际开工日期。')
table(['阶段','主要任务','负责人'],[
('第1周：需求与设计','确认双模式边界；完善原型与数据设计；摄像头两动作可行性试验；建立资料来源清单。交付设计文档与可行性记录。','全体；组员3负责视觉验证'),('第2周：基础业务','实现登录、档案、双模式入口、动作/食物目录、数据库迁移和接口约定。交付可操作页面及基础接口。','组员1、组员2；组员3整理数据'),('第3周：训练核心','完成大众模板、运动员跑步计划、锁定与版本、视觉计数和手工记录。交付两个模式训练演示。','组员1负责页面；组员2计划；组员3视觉'),('第4周：分析与联动','完成饮食结构、负荷汇总、周报及调整预览；处理缺失与冲突。交付完整业务流程。','组员2、组员3；组员1集成'),('第5周：测试与修正','执行隔离测试集、性能/权限/异常测试与五人可用性试用；核查安全边界和来源。修复高优先级缺陷。','全体，交叉复核'),('第6周：交付与验收','冻结必须功能，完成部署恢复演练、报告与PPT、录制演示；有余力再加入可选扩展。','全体；组员2负责部署')],[3.3,10.0,3.1])
para('进度风险控制：第1周若视觉性能不足，优先简化机位与动作配置，并在第3周前完成两个限定动作验证；不可用手工计数冒充视觉验收。第4周后新增功能必须以不影响F01～F12为前提。任何必须功能范围变更均更新需求与验收清单后确认。')
para('关键验收：F01/F02用双账号隔离及切换回归验证；F03/F04/F07用20组约束和并发版本案例验证；F05按3.4.3节隔离样本核对计数；F06/F08用已知输入核对公式和缺失处理；F09/F10用20组餐次场景核对规则；F11检查图表与导出一致；F12完成规则发布、备份恢复和数据删除演练。')

h('9. 预期成果',1,True)
para('（1）可运行的软件系统：同时具有大众模式与运动员模式，完整实现F01～F12；提供本地演示账号与初始化方式，不含真实用户隐私数据。')
para('（2）完整源代码：包含前端、后端、数据库迁移、固定依赖、模型资源获取说明和配置示例，不包含密钥或账号密码。')
para('（3）示例与验证数据：至少20个基础动作条目、40个食物条目、两种模式示例计划、经授权的动作验证片段及标注说明；附数据来源与许可信息。')
para('（4）项目文档：本需求分析与概要设计、后续详细设计与实现说明、测试报告、部署与使用手册、数据与隐私处理说明。')
para('（5）展示材料：项目总结PPT和演示视频，覆盖模式切换、两动作计数、运动员计划实际对照、负荷缺失处理、饮食结构及调整确认。')
para('（6）效果证据：提交计数误差、接口耗时、规则案例和可用性试用记录，同时披露未达到的指标与已知限制。项目成果以软件与验证结果为准，不以减重、竞技成绩提升或医疗效果作为验收承诺。')

h('参考资料',1)
para('资料用于技术选型、规则来源和产品调研。正式实现应固定实际采用的包、模型与资料版本，并在食物、动作及规则条目中记录具体来源。')
refs=['[1] Garmin. Training Load. Garmin Technology. https://www.garmin.com/en-US/garmin-technology/running-science/physiological-measurements/training-load/ 。', '[2] MyFitnessPal. 官方产品与饮食记录功能介绍. https://www.myfitnesspal.com/ 。', '[3] Google AI Edge. Pose Landmarker for Web. https://ai.google.dev/edge/mediapipe/solutions/vision/pose_landmarker/web_js （在线技术指南；实现前复核具体版本）。', '[4] 中国营养学会. 中国居民膳食指南（2022）. 北京：人民卫生出版社，2022. 八准则：http://dg.cnsoc.org/newslist_0402_1.htm 。', '[5] World Health Organization. Physical activity. https://www.who.int/news-room/fact-sheets/detail/physical-activity 。', '[6] Foster C, Florhaug JA, Franklin J, et al. A New Approach to Monitoring Exercise Training. Journal of Strength and Conditioning Research, 2001, 15(1):109–115. PMID:11708692.', '[7] Vue.js. Introduction. https://vuejs.org/guide/introduction.html 。', '[8] FastAPI. 官方文档. https://fastapi.tiangolo.com/ 。']
for r in refs:para(r)

# Preserve original package parts; transplant only changed body, relationships, content types and new images.
tmp=ROOT/'qa'/'working.docx';doc.save(tmp)
with ZipFile(REF) as old, ZipFile(tmp) as new, ZipFile(OUT,'w',ZIP_DEFLATED) as dst:
 change={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'}
 for name in old.namelist(): dst.writestr(name,new.read(name) if name in change else old.read(name))
 for name in new.namelist():
  if name not in old.namelist():dst.writestr(name,new.read(name))
with ZipFile(OUT) as z:
 altered=[n for n,v in inventory.items() if n not in {'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'} and hashlib.sha256(z.read(n)).hexdigest()!=v]
 assert not altered,altered
final=Document(OUT)
body='\n'.join(p.text for p in final.paragraphs)
assert 'XXX' not in body and '建议回答以下问题' not in body
assert len(final.inline_shapes)==8,len(final.inline_shapes)
report={'output':str(OUT),'paragraphs':len(final.paragraphs),'tables':len(final.tables),'figures':len(final.inline_shapes),'preserve_only_changed':altered,'template_sha256':hashlib.sha256(REF.read_bytes()).hexdigest(),'headings':[p.text for p in final.paragraphs if p.style.name.startswith('Course Heading')]}
(ROOT/'qa'/'structure-check.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in report.items() if k!='headings'},ensure_ascii=False,indent=2))
